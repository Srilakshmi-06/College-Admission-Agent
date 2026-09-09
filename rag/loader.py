"""
Document loader — supports PDF, DOCX, and TXT files.
Returns a list of raw document records with extracted text and metadata.
"""

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class RawDocument:
    """A raw document extracted from a file."""
    file_path: str
    file_name: str
    doc_type: str
    pages: list[dict]          # list of {"page": int, "text": str}
    total_pages: int
    error: Optional[str] = None

    @property
    def success(self) -> bool:
        return self.error is None and bool(self.pages)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p["text"] for p in self.pages)


class DocumentLoader:
    """
    Loads documents from PDF, DOCX, or TXT files and extracts text.
    """

    SUPPORTED = {".pdf", ".txt", ".docx"}

    def load_file(self, file_path: str) -> RawDocument:
        """
        Load a single file and return a RawDocument.

        Args:
            file_path: Absolute or relative path to the file.

        Returns:
            RawDocument with extracted pages and metadata.
        """
        path = Path(file_path)
        ext = path.suffix.lower()
        doc_name = path.name

        if not path.exists():
            return RawDocument(
                file_path=str(path),
                file_name=doc_name,
                doc_type=ext.lstrip("."),
                pages=[],
                total_pages=0,
                error=f"File not found: {file_path}",
            )

        if ext not in self.SUPPORTED:
            return RawDocument(
                file_path=str(path),
                file_name=doc_name,
                doc_type=ext.lstrip("."),
                pages=[],
                total_pages=0,
                error=f"Unsupported file type '{ext}'. Supported: {self.SUPPORTED}",
            )

        try:
            if ext == ".pdf":
                return self._load_pdf(path)
            elif ext == ".txt":
                return self._load_txt(path)
            elif ext == ".docx":
                return self._load_docx(path)
        except Exception as e:
            logger.exception(f"Failed to load {file_path}")
            return RawDocument(
                file_path=str(path),
                file_name=doc_name,
                doc_type=ext.lstrip("."),
                pages=[],
                total_pages=0,
                error=f"Load error: {e}",
            )

    def load_directory(self, dir_path: str) -> list[RawDocument]:
        """Load all supported files from a directory."""
        results: list[RawDocument] = []
        dir_ = Path(dir_path)
        if not dir_.exists():
            logger.warning(f"Data directory not found: {dir_path}")
            return results

        for file in sorted(dir_.iterdir()):
            if file.is_file() and file.suffix.lower() in self.SUPPORTED:
                logger.info(f"Loading: {file.name}")
                doc = self.load_file(str(file))
                if doc.success:
                    logger.info(f"  → {doc.total_pages} page(s), {len(doc.full_text)} chars")
                else:
                    logger.warning(f"  → Error: {doc.error}")
                results.append(doc)

        return results

    # ------------------------------------------------------------------
    # Private loaders
    # ------------------------------------------------------------------

    def _load_pdf(self, path: Path) -> RawDocument:
        """Extract text from a PDF using pypdf."""
        try:
            import pypdf  # type: ignore
        except ImportError:
            raise ImportError("pypdf not installed. Run: pip install pypdf")

        pages: list[dict] = []
        try:
            reader = pypdf.PdfReader(str(path))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                text = self._clean_text(text)
                if text.strip():
                    pages.append({"page": i + 1, "text": text})
        except pypdf.errors.PdfReadError as e:
            return RawDocument(
                file_path=str(path),
                file_name=path.name,
                doc_type="pdf",
                pages=[],
                total_pages=0,
                error=f"PDF read error: {e}",
            )

        return RawDocument(
            file_path=str(path),
            file_name=path.name,
            doc_type="pdf",
            pages=pages,
            total_pages=len(reader.pages),
        )

    def _load_txt(self, path: Path) -> RawDocument:
        """Read a plain text file."""
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            return RawDocument(
                file_path=str(path),
                file_name=path.name,
                doc_type="txt",
                pages=[],
                total_pages=0,
                error=f"Text read error: {e}",
            )

        text = self._clean_text(text)
        pages = [{"page": 1, "text": text}] if text.strip() else []
        return RawDocument(
            file_path=str(path),
            file_name=path.name,
            doc_type="txt",
            pages=pages,
            total_pages=1,
        )

    def _load_docx(self, path: Path) -> RawDocument:
        """Extract text from a DOCX file using python-docx."""
        try:
            import docx  # type: ignore
        except ImportError:
            raise ImportError("python-docx not installed. Run: pip install python-docx")

        try:
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            text = self._clean_text(text)
            pages = [{"page": 1, "text": text}] if text.strip() else []
            return RawDocument(
                file_path=str(path),
                file_name=path.name,
                doc_type="docx",
                pages=pages,
                total_pages=1,
            )
        except Exception as e:
            return RawDocument(
                file_path=str(path),
                file_name=path.name,
                doc_type="docx",
                pages=[],
                total_pages=0,
                error=f"DOCX read error: {e}",
            )

    @staticmethod
    def _clean_text(text: str) -> str:
        """Basic text cleaning: normalise whitespace."""
        import re
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\r", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()
